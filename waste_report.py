"""Power BI-style PDF and PowerPoint reports for the waste collection and recovery analysis."""
from __future__ import annotations

import io
import textwrap

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates  # noqa: F401
import matplotlib.pyplot as plt
import matplotlib.ticker as mticker
import pandas as pd
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


# ---------------- Power BI theme ----------------
BLUE, NAVY, RED, GREEN = "#118DFF", "#12239E", "#D64550", "#1AAB40"
ORANGE, GREY, CANVAS, INK, MUTED, GRID = "#E66C37", "#E6E6E6", "#F3F2F1", "#252423", "#605E5C", "#EDEBE9"
SUBTEXT, LABEL = "#E3F1FF", "#CFE6FF"
plt.rcParams.update({
    "font.family": "DejaVu Sans", "font.size": 10, "text.color": INK, "axes.edgecolor": GRID,
    "axes.labelcolor": MUTED, "xtick.color": MUTED, "ytick.color": MUTED,
    "axes.spines.top": False, "axes.spines.right": False, "figure.facecolor": "white",
})
PAGE_W, PAGE_H = landscape(A4)
MARGIN = 28
ALL_SECTIONS = ["summary", "notes", "table", "trends", "areas", "recycling", "vehicles"]




def _clean(ax, grid="y"):
    ax.spines["left"].set_visible(False)
    ax.spines["bottom"].set_color(GRID)
    ax.tick_params(length=0)
    (ax.yaxis if grid == "y" else ax.xaxis).grid(True, color=GRID, linewidth=0.8)
    ax.set_axisbelow(True)


def _t(kg):
    return f"{kg / 1000:,.1f} t"


def _short_month(lbl):
    return lbl.replace(" 20", " '")


MAT_COLOURS = {"PET bottles": BLUE, "Plastic": "#744EC2", "Metal & cans": NAVY, "Glass": "#1AAB40",
               "Beer bottles": "#D9B300", "Carton & paper": ORANGE, "Others": "#A19F9D"}


def chart_pages(data: dict, sections=None):
    """Yield (title, subtitle, figure) for each chart page – shared by the PDF and PowerPoint.
    data = dict(monthly, areas, vehicles, materials, daily)."""
    secs = set(sections or ALL_SECTIONS)
    m, areas, veh, mats, daily = (data[k] for k in ("monthly", "areas", "vehicles", "materials", "daily"))
    months = [_short_month(x) for x in m.index]

    # ---- Monthly table ----
    if "table" in secs:
        rows = [("Waste collected (kg)", "Waste collected (kg)", "{:,.0f}"),
                ("Collection trips", "Collection trips", "{:,.0f}"),
                ("Collection days", "Collection days", "{:,.0f}"),
                ("Avg per collection day (kg)", "Avg per collection day (kg)", "{:,.0f}"),
                ("Avg per calendar day (kg)", "Avg per calendar day (kg)", "{:,.0f}"),
                ("Recyclables recovered (kg)", "Recyclables recovered (kg)", "{:,.0f}"),
                ("Food waste recovered (kg)", "Food waste recovered (kg)", "{:,.0f}"),
                ("Industrial waste (kg)", "Industrial waste (kg)", "{:,.0f}"),
                ("Total waste generated (kg)", "Total waste generated (kg)", "{:,.0f}"),
                ("Generated per day (kg)", "Generated per day (kg)", "{:,.0f}"),
                ("Recovery %", "Recovery %", "{:.1f}%")]
        cell = [[("no data" if pd.isna(m.loc[mo, col]) else fmt.format(m.loc[mo, col])) for mo in m.index]
                for col, _, fmt in rows]
        labels = [f"{_short_month(x)}{' *' if m.loc[x, 'Partial month'] else ''}" for x in m.index]
        fig, ax = plt.subplots(figsize=(13, 0.48 * len(rows) + 1.2))
        ax.axis("off")
        t = ax.table(cellText=cell, rowLabels=[r[1] for r in rows], colLabels=labels, loc="center",
                     cellLoc="center", rowLoc="left")
        t.auto_set_font_size(False)
        t.set_fontsize(10)
        t.scale(1, 1.7)
        for (r, c), cl in t.get_celld().items():
            cl.set_edgecolor(GRID)
            if r == 0:
                cl.set_facecolor(BLUE)
                cl.get_text().set_color("white")
                cl.get_text().set_fontweight("bold")
            if c == -1:
                cl.get_text().set_fontweight("bold")
                cl.set_facecolor("#F7F7F7")
            if r == len(rows):
                cl.set_facecolor("#E3F5E8" if c != -1 else "#F7F7F7")
                cl.get_text().set_fontweight("bold")
        if m["Partial month"].any():
            fig.text(0.01, 0.01, "* Partial month – collection records start part-way through the month.",
                     fontsize=9, color=MUTED)
        fig.tight_layout()
        yield "Monthly summary", "Waste collected, recovered and generated each month", fig

    # ---- Trends ----
    if "trends" in secs:
        fig, axes = plt.subplots(1, 2, figsize=(15, 4.9), gridspec_kw={"width_ratios": [1.2, 1]})
        x = range(len(months))
        w = 0.38
        ax = axes[0]
        ax.bar([i - w / 2 for i in x], m["Waste collected (kg)"] / 1000, w, color=BLUE, label="Collected (to landfill)")
        ax.bar([i + w / 2 for i in x], m["Total recovered (kg)"] / 1000, w, color=GREEN, label="Recovered (recycled + food)")
        for i, (a, b) in enumerate(zip(m["Waste collected (kg)"], m["Total recovered (kg)"])):
            ax.text(i - w / 2, a / 1000 + 2, f"{a / 1000:,.0f}", ha="center", fontsize=8.5)
            ax.text(i + w / 2, b / 1000 + 2, f"{b / 1000:,.0f}", ha="center", fontsize=8.5)
        ax.set_xticks(list(x), months)
        ax.set_ylabel("Tonnes")
        ax.set_title("Collected vs recovered (tonnes)\n", loc="left", fontsize=11.5, fontweight="bold")
        ax.legend(loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2, frameon=False, fontsize=9)
        _clean(ax)
        ax2 = axes[1]
        full = m[m["Recycling data"]]
        ax2.plot([_short_month(i) for i in full.index], full["Recovery %"], color=GREEN, marker="o", linewidth=2.5)
        for i, v in enumerate(full["Recovery %"]):
            ax2.annotate(f"{v:.1f}%", (i, v), textcoords="offset points", xytext=(0, 9), ha="center",
                         fontsize=10, fontweight="bold", color=GREEN)
        ax2.set_ylim(0, max(60, full["Recovery %"].max() * 1.3 if len(full) else 60))
        ax2.yaxis.set_major_formatter(mticker.PercentFormatter(decimals=0))
        ax2.set_title("Recovery rate\n", loc="left", fontsize=11.5, fontweight="bold")
        ax2.text(0, 1.02, "Recovered ÷ total waste generated", transform=ax2.transAxes, fontsize=8.5, color=MUTED)
        _clean(ax2)
        fig.tight_layout(w_pad=3)
        yield "Collection and recovery trends", "Monthly tonnes and recovery rate", fig

        if daily is not None and len(daily):
            fig, ax = plt.subplots(figsize=(15, 4.3))
            ax.bar(daily.index, daily.values / 1000, color=BLUE, width=0.8)
            avg = daily.mean() / 1000
            ax.axhline(avg, color=ORANGE, linestyle="--", linewidth=1.3)
            ax.text(daily.index[0], avg * 1.04, f"Average {avg:,.1f} t per collection day", color=ORANGE, fontsize=9)
            ax.set_ylabel("Tonnes")
            ax.set_title("Waste collected per day\n", loc="left", fontsize=11.5, fontweight="bold")
            ax.xaxis.set_major_formatter(matplotlib.dates.DateFormatter("%d %b"))
            _clean(ax)
            fig.tight_layout()
            yield "Daily collection", "Total weighed each collection day", fig

    # ---- Areas ----
    if "areas" in secs and not areas.empty:
        a = areas.head(12)
        fig, ax = plt.subplots(figsize=(15, 0.45 * len(a) + 1.4))
        names = ["\n".join(textwrap.wrap(n, 38)) for n in a.index][::-1]
        vals = (a["Waste collected (kg)"] / 1000).values[::-1]
        ax.barh(names, vals, color=BLUE, height=0.62)
        for i, (v, s, n, f) in enumerate(zip(vals, a["Share %"].values[::-1], a["Collections"].values[::-1],
                                              a["Collections per week"].values[::-1])):
            ax.text(v + vals.max() * 0.01, i, f"{v:,.1f} t  ({s:.0f}%) · {int(n)} trips · {f:.1f}/week",
                    va="center", fontsize=9)
        ax.set_xlim(0, vals.max() * 1.45)
        ax.set_xlabel("Tonnes")
        ax.tick_params(axis="y", labelsize=9)
        _clean(ax, grid="x")
        fig.tight_layout()
        yield "Waste collected by area", "Tonnes, share of total, trips and collections per week", fig

    # ---- Recyclables ----
    if "recycling" in secs and mats is not None and not mats.empty:
        fig, axes = plt.subplots(1, 2, figsize=(15, 5.2), gridspec_kw={"width_ratios": [1.45, 1]})
        ax = axes[0]
        src = mats.index[::-1]
        left = [0] * len(src)
        for mat in [c for c in mats.columns if c != "Total (kg)"]:
            vals = (mats.loc[src, mat] / 1000).values
            ax.barh([s[:28] for s in src], vals, left=left, color=MAT_COLOURS.get(mat, GREY), label=mat, height=0.62)
            left = [l + v for l, v in zip(left, vals)]
        for i, tot in enumerate(left):
            ax.text(tot + max(left) * 0.01, i, f"{tot:,.1f} t", va="center", fontsize=9, fontweight="bold")
        ax.set_xlim(0, max(left) * 1.18)
        ax.set_xlabel("Tonnes")
        ax.legend(loc="lower left", bbox_to_anchor=(0, 1.0), ncol=4, frameon=False, fontsize=8.5)
        ax.set_title("Recyclables by collector\n\n", loc="left", fontsize=11.5, fontweight="bold")
        _clean(ax, grid="x")
        ax2 = axes[1]
        tot = mats.drop(columns="Total (kg)").sum().sort_values(ascending=False)
        tot = tot[tot > 0]
        wedges, _ = ax2.pie(tot.values, colors=[MAT_COLOURS.get(k, GREY) for k in tot.index], startangle=90,
                            counterclock=False, wedgeprops=dict(width=0.38, edgecolor="white", linewidth=1.5))
        ax2.text(0, 0.05, f"{tot.sum() / 1000:,.1f} t", ha="center", fontsize=15, fontweight="bold")
        ax2.text(0, -0.18, "recovered", ha="center", fontsize=8.5, color=MUTED)
        ax2.legend(wedges, [f"{k}  {v / 1000:,.1f} t ({v / tot.sum() * 100:.0f}%)" for k, v in tot.items()],
                   loc="center left", bbox_to_anchor=(0.95, 0.5), frameon=False, fontsize=9)
        ax2.set_title("By material", loc="left", fontsize=11.5, fontweight="bold")
        fig.tight_layout(w_pad=2)
        yield "Recyclables recovered", "RVM and scrap dealers – from the daily sheets", fig

    # ---- Vehicles ----
    if "vehicles" in secs and veh is not None and not veh.empty:
        fig, ax = plt.subplots(figsize=(15, 0.5 * len(veh) + 1.4))
        names = list(veh.index)[::-1]
        vals = (veh["Waste carried (kg)"] / 1000).values[::-1]
        ax.barh(names, vals, color=NAVY, height=0.6)
        for i, (v, t_, a_, ar) in enumerate(zip(vals, veh["Trips"].values[::-1], veh["Avg per trip (kg)"].values[::-1],
                                                veh["Areas served"].values[::-1])):
            ax.text(v + vals.max() * 0.01, i, f"{v:,.1f} t · {int(t_)} trips · {a_:,.0f} kg/trip · {int(ar)} areas",
                    va="center", fontsize=9)
        ax.set_xlim(0, vals.max() * 1.5)
        ax.set_xlabel("Tonnes")
        _clean(ax, grid="x")
        fig.tight_layout()
        yield "Vehicles", "Waste carried, trips and average load per trip", fig


# ---------------- KPI cards ----------------

def kpi_cards(data: dict, month: str):
    r = data["monthly"].loc[month]
    star = " (partial month)" if r["Partial month"] else ""
    return [
        (_t(r["Waste collected (kg)"]), f"Waste collected{star}", None),
        (f"{int(r['Collection trips'])}", f"Collection trips on {int(r['Collection days'])} days", None),
        (f"{r['Avg per calendar day (kg)']:,.0f} kg", "Collected per calendar day", None),
        (_t(r["Total waste generated (kg)"]), "Total waste generated (collected + recovered + industrial)", None),
        (_t(r["Recyclables recovered (kg)"]), "Recyclables recovered (RVM + scrap dealers)", None),
        (_t(r["Food waste recovered (kg)"]), "Food waste recovered", None),
        ("no data" if pd.isna(r["Recovery %"]) else f"{r['Recovery %']:.1f}%",
         "Recovery rate (recovered ÷ generated)", None),
        (f"{r['Generated per day (kg)']:,.0f} kg", "Waste generated per day", None),
    ]


# ---------------- PDF ----------------

def _wrap(text, width, max_lines):
    lines = textwrap.wrap(str(text), max(10, width))
    if len(lines) > max_lines:
        lines = lines[:max_lines]
        lines[-1] = lines[-1][:-1] + "…"
    return lines


def _bg(c):
    c.setFillColor(CANVAS)
    c.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)


def _header(c, title, subtitle):
    c.setFillColor(BLUE)
    c.rect(0, PAGE_H - 62, PAGE_W, 62, fill=1, stroke=0)
    c.setFillColor(NAVY)
    c.rect(0, PAGE_H - 62, 6, 62, fill=1, stroke=0)
    c.setFillColor("white")
    c.setFont("Helvetica-Bold", 20)
    c.drawString(MARGIN, PAGE_H - 34, title)
    if subtitle:
        c.setFillColor(SUBTEXT)
        c.setFont("Helvetica", 10.5)
        c.drawString(MARGIN, PAGE_H - 51, subtitle)


def _footer(c, text):
    c.setFillColor(MUTED)
    c.setFont("Helvetica", 8.5)
    c.drawString(MARGIN, 14, text)
    c.drawRightString(PAGE_W - MARGIN, 14, f"Page {c.getPageNumber()}")


def _card(c, x, y, w, h, top=None):
    c.setFillColor("#E1DFDD")
    c.roundRect(x + 1, y - 1.5, w, h, 5, fill=1, stroke=0)
    c.setFillColor("white")
    c.roundRect(x, y, w, h, 5, fill=1, stroke=0)
    if top:
        c.setFillColor(top)
        c.rect(x + 2, y + h - 4, w - 4, 4, fill=1, stroke=0)


def _fig_img(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return ImageReader(buf)


def _cover(c, title, role, author, period_line):
    c.setFillColor(BLUE)
    c.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)
    c.saveState()
    c.setFillColor("white")
    for cx, cy, r, a in [(PAGE_W - 30, PAGE_H - 20, 240, 0.07), (PAGE_W - 120, 40, 170, 0.06),
                         (PAGE_W - 250, PAGE_H / 2 - 10, 90, 0.05)]:
        c.setFillAlpha(a)
        c.circle(cx, cy, r, fill=1, stroke=0)
    c.restoreState()
    c.setFillColor(NAVY)
    c.rect(0, 0, 16, PAGE_H, fill=1, stroke=0)
    x0 = 64
    c.setFillColor(LABEL)
    c.setFont("Helvetica-Bold", 10.5)
    c.drawString(x0, PAGE_H - 70, "R E P O R T")
    c.setFillColor("white")
    c.rect(x0, PAGE_H - 92, 64, 4, fill=1, stroke=0)
    size = 40 if len(title) <= 36 else 34
    y = PAGE_H - 92 - 22 - size
    c.setFont("Helvetica-Bold", size)
    for row in _wrap(title, 32 if size == 40 else 38, 3):
        c.drawString(x0, y, row)
        y -= size + 8
    c.setFillColor(SUBTEXT)
    c.setFont("Helvetica", 13)
    c.drawString(x0, y - 4, period_line[:120])
    base = 70
    c.setStrokeColor("white")
    c.setStrokeAlpha(0.55)
    c.setLineWidth(0.8)
    c.line(x0, base + 92, x0 + 300, base + 92)
    c.setStrokeAlpha(1)
    if author or role:
        c.setFillColor(LABEL)
        c.setFont("Helvetica-Bold", 9.5)
        c.drawString(x0, base + 70, "P R E P A R E D   B Y")
    if author:
        c.setFillColor("white")
        c.setFont("Helvetica-Bold", 22)
        c.drawString(x0, base + 40, author)
    if role:
        c.setFillColor(SUBTEXT)
        c.setFont("Helvetica", 14)
        c.drawString(x0, base + (16 if author else 40), role)


def build_pdf(data, month, cover, notes=None, sections=None, data_notes=None) -> bytes:
    """cover = dict(title, role, author, period_line)."""
    secs = set(sections or ALL_SECTIONS)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=landscape(A4))
    footer = " – ".join(x for x in (cover.get("title"), cover.get("role")) if x)
    _cover(c, cover.get("title", ""), cover.get("role", ""), cover.get("author", ""), cover.get("period_line", ""))
    c.showPage()

    if "summary" in secs:
        _bg(c)
        _header(c, f"Summary – {month}", "Key figures for the month")
        cards = kpi_cards(data, month)
        per_row, gap = 4, 12
        kw = (PAGE_W - 2 * MARGIN - gap * (per_row - 1)) / per_row
        kh = 100
        for i, (value, label, ok) in enumerate(cards):
            row, col = divmod(i, per_row)
            x = MARGIN + col * (kw + gap)
            y = PAGE_H - 62 - 18 - kh - row * (kh + gap)
            top = BLUE if ok is None else (GREEN if ok else RED)
            _card(c, x, y, kw, kh, top)
            c.setFillColor(INK if ok is None else (GREEN if ok else RED))
            c.setFont("Helvetica-Bold", 30)
            c.drawString(x + 16, y + kh - 46, value)
            c.setFillColor(MUTED)
            c.setFont("Helvetica", 10)
            lines = [ln for part in label.split("  ·  ") for ln in _wrap(part, int(kw / 5.2), 2)][:3]
            for j, line in enumerate(lines):
                c.drawString(x + 16, y + kh - 64 - j * 11.5, line)
            if ok is not None:
                c.setFont("Helvetica-Bold", 9)
                c.setFillColor(GREEN if ok else RED)
                c.drawRightString(x + kw - 12, y + kh - 22, "MET" if ok else "NOT MET")
        if data_notes:
            y0 = PAGE_H - 62 - 18 - 2 * (kh + gap) - 14
            c.setFillColor(MUTED)
            c.setFont("Helvetica-Oblique", 9)
            for i, n in enumerate(data_notes[:4]):
                c.drawString(MARGIN, y0 - i * 12, f"Note: {n}"[:170])
        _footer(c, footer)
        c.showPage()

    if "notes" in secs and notes and (notes.get("findings") or notes.get("recommendations")):
        _bg(c)
        _header(c, "Key findings & recommendations", cover.get("period_line", ""))
        boxes = [(h, it) for h, it in (("Key findings", notes.get("findings") or []),
                                       ("Recommendations", notes.get("recommendations") or [])) if it]
        gap = 12
        bw = (PAGE_W - 2 * MARGIN - gap * (len(boxes) - 1)) / len(boxes)
        by, bh = 32, PAGE_H - 62 - 16 - 32
        for bi, (heading, items) in enumerate(boxes):
            x = MARGIN + bi * (bw + gap)
            _card(c, x, by, bw, bh, BLUE if bi == 0 else NAVY)
            c.setFillColor(INK)
            c.setFont("Helvetica-Bold", 15)
            c.drawString(x + 18, by + bh - 34, heading)
            yy = by + bh - 62
            chars = int((bw - 60) / 6.4)
            for n, item in enumerate(items, 1):
                lines = textwrap.wrap(item, chars) or [""]
                if yy - 18 * len(lines) < by + 14:
                    break
                c.setFillColor(BLUE if bi == 0 else NAVY)
                c.setFont("Helvetica-Bold", 13)
                c.drawString(x + 18, yy, f"{n}." if bi == 1 else "•")
                c.setFillColor(INK)
                c.setFont("Helvetica", 13)
                for line in lines:
                    c.drawString(x + 38, yy, line)
                    yy -= 18
                yy -= 9
        _footer(c, footer)
        c.showPage()

    for title, subtitle, fig in chart_pages(data, secs):
        _bg(c)
        _header(c, title, subtitle)
        img = _fig_img(fig)
        x, y, w, h = MARGIN, 32, PAGE_W - 2 * MARGIN, PAGE_H - 62 - 16 - 32
        _card(c, x, y, w, h)
        iw, ih = img.getSize()
        s = min((w - 28) / iw, (h - 28) / ih)
        c.drawImage(img, x + (w - iw * s) / 2, y + h - 14 - ih * s, iw * s, ih * s)
        _footer(c, footer)
        c.showPage()
    c.save()
    return buf.getvalue()


# ---------------- PowerPoint ----------------

def build_pptx(data, month, cover, notes=None, sections=None, data_notes=None) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    def rgb(h):
        h = h.lstrip("#")
        return RGBColor(int(h[:2], 16), int(h[2:4], 16), int(h[4:], 16))

    secs = set(sections or ALL_SECTIONS)
    W, H, HH = Inches(13.333), Inches(7.5), Inches(0.95)
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    footer = " – ".join(x for x in (cover.get("title"), cover.get("role")) if x)

    def rect(sl, x, y, w, h, col, shape=MSO_SHAPE.RECTANGLE):
        s = sl.shapes.add_shape(shape, x, y, w, h)
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0] = 0.04
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(col)
        s.line.fill.background()
        s.shadow.inherit = False
        return s

    def text(sl, t, x, y, w, h, size=14, bold=False, col=INK):
        tb = sl.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(str(t).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size, p.font.bold, p.font.name = Pt(size), bold, "Segoe UI"
            p.font.color.rgb = rgb(col)
            p.space_after = Pt(6)

    def card(sl, x, y, w, h, top=None):
        rect(sl, x + Emu(12000), y + Emu(15000), w, h, "#E1DFDD", MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(sl, x, y, w, h, "#FFFFFF", MSO_SHAPE.ROUNDED_RECTANGLE)
        if top:
            rect(sl, x + Inches(0.05), y, w - Inches(0.1), Inches(0.06), top)

    def base(title, subtitle):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.background.fill.solid()
        sl.background.fill.fore_color.rgb = rgb(CANVAS)
        rect(sl, 0, 0, W, HH, BLUE)
        rect(sl, 0, 0, Inches(0.09), HH, NAVY)
        text(sl, title, Inches(0.4), Inches(0.14), Inches(12), Inches(0.5), 24, True, "#FFFFFF")
        if subtitle:
            text(sl, subtitle, Inches(0.4), Inches(0.58), Inches(12), Inches(0.3), 12, False, SUBTEXT)
        text(sl, footer, Inches(0.4), Inches(7.15), Inches(11), Inches(0.25), 9, False, MUTED)
        return sl

    # Cover
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = rgb(BLUE)
    for cx, cy, r, col in [(13.0, 0.2, 3.3, "#2A99FF"), (11.9, 7.3, 2.4, "#2696FF"), (9.7, 3.6, 1.25, "#2A99FF")]:
        rect(sl, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r), col, MSO_SHAPE.OVAL)
    rect(sl, 0, 0, Inches(0.22), H, NAVY)
    text(sl, "R E P O R T", Inches(0.9), Inches(0.75), Inches(4), Inches(0.3), 11, True, LABEL)
    rect(sl, Inches(0.9), Inches(1.15), Inches(0.9), Inches(0.06), "#FFFFFF")
    text(sl, cover.get("title", ""), Inches(0.9), Inches(1.4), Inches(8.4), Inches(1.9), 44, True, "#FFFFFF")
    text(sl, cover.get("period_line", ""), Inches(0.9), Inches(3.35), Inches(9), Inches(0.4), 14, False, SUBTEXT)
    rect(sl, Inches(0.9), Inches(5.35), Inches(4.2), Emu(9525), "#FFFFFF")
    if cover.get("author") or cover.get("role"):
        text(sl, "P R E P A R E D   B Y", Inches(0.9), Inches(5.55), Inches(5), Inches(0.3), 10, True, LABEL)
    if cover.get("author"):
        text(sl, cover["author"], Inches(0.9), Inches(5.85), Inches(8), Inches(0.5), 26, True, "#FFFFFF")
    if cover.get("role"):
        text(sl, cover["role"], Inches(0.9), Inches(6.45 if cover.get("author") else 5.85), Inches(8), Inches(0.4),
             16, False, SUBTEXT)

    if "summary" in secs:
        sl = base(f"Summary – {month}", "Key figures for the month")
        cards = kpi_cards(data, month)
        left, gap = Inches(0.35), Inches(0.2)
        kw = int((W - 2 * left - gap * 3) / 4)
        kh = Inches(1.65)
        for i, (value, label, ok) in enumerate(cards):
            row, col = divmod(i, 4)
            x = left + col * (kw + gap)
            y = HH + Inches(0.25) + row * (kh + gap)
            card(sl, x, y, kw, kh, BLUE if ok is None else (GREEN if ok else RED))
            text(sl, value, x + Inches(0.2), y + Inches(0.2), kw - Inches(0.3), Inches(0.6), 30, True,
                 INK if ok is None else (GREEN if ok else RED))
            text(sl, label, x + Inches(0.2), y + Inches(0.9), kw - Inches(0.35), Inches(0.6), 11, False, MUTED)
            if ok is not None:
                text(sl, "MET" if ok else "NOT MET", x + kw - Inches(1.2), y + Inches(0.18), Inches(1.0), Inches(0.3),
                     10, True, GREEN if ok else RED)
        if data_notes:
            text(sl, "\n".join("Note: " + n for n in data_notes[:4]), Inches(0.4), Inches(5.5), Inches(12.4),
                 Inches(1.4), 10, False, MUTED)

    if "notes" in secs and notes and (notes.get("findings") or notes.get("recommendations")):
        sl = base("Key findings & recommendations", cover.get("period_line", ""))
        boxes = [(h, it) for h, it in (("Key findings", notes.get("findings") or []),
                                       ("Recommendations", notes.get("recommendations") or [])) if it]
        left, gap = Inches(0.35), Inches(0.2)
        bw = int((W - 2 * left - gap * (len(boxes) - 1)) / len(boxes))
        by = HH + Inches(0.2)
        bh = H - by - Inches(0.45)
        for bi, (heading, items) in enumerate(boxes):
            x = left + bi * (bw + gap)
            card(sl, x, by, bw, bh, BLUE if bi == 0 else NAVY)
            text(sl, heading, x + Inches(0.3), by + Inches(0.25), bw - Inches(0.6), Inches(0.4), 18, True)
            body = "\n".join((f"{n}.  " if bi == 1 else "•  ") + it for n, it in enumerate(items, 1))
            text(sl, body, x + Inches(0.3), by + Inches(0.8), bw - Inches(0.6), bh - Inches(1.0),
                 14 if sum(map(len, items)) < 700 else 12)

    for title, subtitle, fig in chart_pages(data, secs):
        sl = base(title, subtitle)
        x, y = Inches(0.35), HH + Inches(0.2)
        w, h = W - Inches(0.7), H - y - Inches(0.45)
        card(sl, x, y, w, h)
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=200, bbox_inches="tight", facecolor="white")
        fw, fh = fig.get_size_inches()
        plt.close(fig)
        buf.seek(0)
        pad = Inches(0.2)
        s = min((w - 2 * pad) / fw, (h - 2 * pad) / fh)
        pw, ph = Emu(int(fw * s)), Emu(int(fh * s))
        sl.shapes.add_picture(buf, int(x + (w - pw) / 2), int(y + pad), pw, ph)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ---------------- Excel ----------------

def build_xlsx(data: dict, cover: dict, data_notes=None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="D0D7E2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill, head_font = PatternFill("solid", fgColor=BLUE[1:]), Font(bold=True, color="FFFFFF")
    wb = Workbook()
    first = True
    for name, df, idx in [("Monthly summary", data["monthly"].drop(columns=["Partial month", "Recycling data"]), "Month"),
                          ("By area", data["areas"], "Area"), ("Recyclables", data["materials"], "Collector"),
                          ("Vehicles", data["vehicles"], "Vehicle")]:
        if df is None or df.empty:
            continue
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws["A1"] = f"{cover.get('title', '')} – {name}"
        ws["A1"].font = Font(bold=True, size=14, color=BLUE[1:])
        ws["A2"] = " · ".join(x for x in (cover.get("role"), cover.get("author"), cover.get("period_line")) if x)
        ws["A2"].font = Font(italic=True, color="605E5C")
        hdr = [idx] + list(df.columns)
        for j, h in enumerate(hdr, 1):
            c = ws.cell(row=4, column=j, value=h)
            c.fill, c.font, c.border = head_fill, head_font, border
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        for i, (k, row) in enumerate(df.iterrows(), start=5):
            ws.cell(row=i, column=1, value=str(k)).border = border
            for j, v in enumerate(row, start=2):
                val = (None if pd.isna(v) else round(float(v), 1)) if isinstance(v, (int, float)) and not isinstance(v, bool) else v
                c = ws.cell(row=i, column=j, value=val)
                c.border = border
                if isinstance(val, float):
                    c.number_format = "#,##0.0"
        ws.column_dimensions["A"].width = 42 if idx in ("Area", "Collector") else 14
        for j in range(2, len(hdr) + 1):
            ws.column_dimensions[ws.cell(row=4, column=j).column_letter].width = 15
        ws.row_dimensions[4].height = 44
        ws.freeze_panes = "B5"
        if name == "Monthly summary":
            r = 6 + len(df)
            notes = ["Recovery % = (recyclables + food waste) ÷ (collected + recyclables + food waste + industrial).",
                     "Recyclables are summed from the daily RVM and scrap-dealer sheets; food and industrial waste "
                     "from the monthly report sheets."] + [f"Note: {n}" for n in (data_notes or [])]
            for k, line in enumerate(notes):
                ws.cell(row=r + k, column=1, value=line).font = Font(italic=True, color="605E5C")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
